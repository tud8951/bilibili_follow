import tkinter as tk
from tkinter import ttk, scrolledtext, filedialog, messagebox
import threading
import requests
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
import os
import time


class BilibiliFollowExport:
    def __init__(self, root):
        self.root = root
        self.root.title("B站关注列表导出工具")
        self.root.geometry("640x560")
        self.root.resizable(False, False)

        # API 相关
        self.session = requests.Session()
        self.session.headers.update({
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                          "AppleWebKit/537.36 (KHTML, like Gecko) "
                          "Chrome/120.0.0.0 Safari/537.36",
            "Referer": "https://www.bilibili.com/",
        })

        self._build_ui()

    def _build_ui(self):
        main_frame = ttk.Frame(self.root, padding=16)
        main_frame.pack(fill=tk.BOTH, expand=True)

        # ========== Cookie 输入 ==========
        ttk.Label(main_frame, text="Cookie (SESSDATA):", font=("Microsoft YaHei", 10)).pack(
            anchor=tk.W, pady=(0, 4)
        )
        cookie_frame = ttk.Frame(main_frame)
        cookie_frame.pack(fill=tk.X, pady=(0, 8))

        self.cookie_var = tk.StringVar()
        self.cookie_entry = ttk.Entry(
            cookie_frame, textvariable=self.cookie_var, show="*", font=("Consolas", 10)
        )
        self.cookie_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=2)

        self.show_cookie_btn = ttk.Button(
            cookie_frame, text="显示", width=6, command=self._toggle_cookie_visibility
        )
        self.show_cookie_btn.pack(side=tk.LEFT, padx=(6, 0))

        self.browser_cookie_btn = ttk.Button(
            cookie_frame, text="网页登录获取", command=self._start_browser_cookie
        )
        self.browser_cookie_btn.pack(side=tk.LEFT, padx=(6, 0))

        # Cookie 获取提示
        tip_label = ttk.Label(
            main_frame,
            text="💡 点击“网页登录获取”，在弹出的 B站窗口中完成登录，程序会自动获取 Cookie 并关闭窗口",
            foreground="gray",
            font=("Microsoft YaHei", 9),
            wraplength=580,
        )
        tip_label.pack(anchor=tk.W, pady=(0, 8))

        # ========== 操作按钮 ==========
        btn_frame = ttk.Frame(main_frame)
        btn_frame.pack(fill=tk.X, pady=(0, 8))

        self.fetch_btn = ttk.Button(
            btn_frame, text="🚀 获取关注列表", command=self._start_fetch
        )
        self.fetch_btn.pack(side=tk.LEFT, padx=(0, 8))

        self.export_btn = ttk.Button(
            btn_frame, text="📥 导出 Excel", command=self._export_excel, state=tk.DISABLED
        )
        self.export_btn.pack(side=tk.LEFT)

        self.progress = ttk.Progressbar(main_frame, mode="indeterminate")
        self.progress.pack(fill=tk.X, pady=(0, 8))

        # ========== 状态信息 ==========
        self.status_var = tk.StringVar(value="就绪")
        status_bar = ttk.Label(
            main_frame, textvariable=self.status_var, font=("Microsoft YaHei", 9)
        )
        status_bar.pack(anchor=tk.W, pady=(0, 4))

        # ========== 日志区域 ==========
        log_label = ttk.Label(main_frame, text="运行日志:", font=("Microsoft YaHei", 10))
        log_label.pack(anchor=tk.W, pady=(0, 4))

        self.log_area = scrolledtext.ScrolledText(
            main_frame,
            height=16,
            font=("Consolas", 9),
            bg="#1e1e1e",
            fg="#d4d4d4",
            insertbackground="white",
            state=tk.DISABLED,
            wrap=tk.WORD,
        )
        self.log_area.pack(fill=tk.BOTH, expand=True)

        # 数据缓存
        self.follow_data = []  # [(uid, name), ...]
        self.follow_total = 0
        self.current_page = 0
        self.my_name = ""  # 当前用户昵称
        self.my_uid = ""  # 当前用户 UID

    # ---------- 工具方法 ----------

    def _toggle_cookie_visibility(self):
        if self.cookie_entry.cget("show") == "*":
            self.cookie_entry.configure(show="")
            self.show_cookie_btn.configure(text="隐藏")
        else:
            self.cookie_entry.configure(show="*")
            self.show_cookie_btn.configure(text="显示")

    def _start_browser_cookie(self):
        self.browser_cookie_btn.configure(state=tk.DISABLED)
        self.status_var.set("请在弹出的 B站窗口中登录…")
        self._log("已打开独立 B站登录窗口，请完成登录…")
        threading.Thread(target=self._login_in_browser, daemon=True).start()

    def _login_in_browser(self):
        try:
            from playwright.sync_api import sync_playwright

            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(headless=False)
                context = browser.new_context()
                page = context.new_page()
                page.goto("https://www.bilibili.com", wait_until="domcontentloaded")

                sessdata = None
                for _ in range(300):
                    if page.is_closed():
                        break
                    cookies = context.cookies("https://www.bilibili.com")
                    sessdata = next(
                        (cookie["value"] for cookie in cookies if cookie["name"] == "SESSDATA"),
                        None,
                    )
                    if sessdata:
                        break
                    page.wait_for_timeout(1000)

                if not sessdata:
                    raise RuntimeError("登录窗口已关闭，或等待登录超时")

                self.root.after(0, lambda: self.cookie_var.set(sessdata))
                self.root.after(0, lambda: self._log("✅ 登录成功，已自动获取 SESSDATA"))
                context.close()
                browser.close()
        except ImportError:
            self.root.after(
                0,
                lambda: messagebox.showerror(
                    "缺少依赖", "请先运行：pip install playwright，然后运行：playwright install chromium"
                ),
            )
        except Exception as e:
            error_message = str(e)
            self.root.after(0, lambda: self._log(f"❌ 获取浏览器 Cookie 失败: {error_message}"))
            self.root.after(0, lambda: messagebox.showerror("获取失败", error_message))
        finally:
            self.root.after(0, lambda: self.browser_cookie_btn.configure(state=tk.NORMAL))
            self.root.after(0, lambda: self.status_var.set("就绪"))

    def _log(self, msg: str):
        self.log_area.configure(state=tk.NORMAL)
        self.log_area.insert(tk.END, msg + "\n")
        self.log_area.see(tk.END)
        self.log_area.configure(state=tk.DISABLED)

    def _set_ui_busy(self, busy: bool):
        state = tk.DISABLED if busy else tk.NORMAL
        self.fetch_btn.configure(state=state)
        self.export_btn.configure(state=state)
        if busy:
            self.progress.start(10)
            self.status_var.set("处理中…")
        else:
            self.progress.stop()
            self.status_var.set("就绪")

    # ---------- 核心逻辑 ----------

    def _start_fetch(self):
        cookie_value = self.cookie_var.get().strip()
        if not cookie_value:
            messagebox.showwarning("提示", "请先输入 SESSDATA Cookie")
            return

        self.follow_data.clear()
        self.follow_total = 0
        self.current_page = 0
        self._log("=" * 50)
        self._log("开始获取关注列表…")
        self._set_ui_busy(True)
        self.export_btn.configure(state=tk.DISABLED)

        threading.Thread(target=self._fetch_all, daemon=True).start()

    def _fetch_all(self):
        try:
            # 设置 cookie
            self.session.cookies.set("SESSDATA", self.cookie_var.get().strip())

            # Step 1: 获取用户自己的信息
            self._log("[1/2] 正在获取用户信息…")
            info_resp = self.session.get("https://api.bilibili.com/x/space/myinfo")
            info_data = info_resp.json()

            if info_data.get("code") != 0:
                err_msg = info_data.get("message", "未知错误")
                self._log(f"❌ 获取用户信息失败: {err_msg}")
                self.root.after(0, lambda: messagebox.showerror("错误", f"获取用户信息失败: {err_msg}"))
                self.root.after(0, lambda: self._set_ui_busy(False))
                return

            my_uid = info_data["data"]["mid"]
            my_name = info_data["data"]["name"]
            self.my_name = my_name
            self.my_uid = str(my_uid)
            self._log(f"   ✅ 当前用户: {my_name} (UID: {my_uid})")

            # Step 2: 分页获取关注列表
            self._log("[2/2] 正在获取关注列表…")
            page = 1
            ps = 50  # 每页 50 条

            while True:
                self._log(f"   正在获取第 {page} 页…")
                url = f"https://api.bilibili.com/x/relation/followings?vmid={my_uid}&pn={page}&ps={ps}&order=desc"
                resp = self.session.get(url)
                data = resp.json()

                if data.get("code") != 0:
                    err_msg = data.get("message", "未知错误")
                    self._log(f"❌ 获取第 {page} 页失败: {err_msg}")
                    break

                page_data = data["data"]
                total = page_data.get("total", 0)
                follow_list = page_data.get("list", [])

                if page == 1:
                    self.follow_total = total
                    self._log(f"   共关注 {total} 人")

                for item in follow_list:
                    uid = item["mid"]
                    name = item["uname"]
                    self.follow_data.append((uid, name))

                self._log(f"   第 {page} 页完成，已累计获取 {len(self.follow_data)} 条")

                # 判断是否还有下一页
                if len(follow_list) < ps or len(self.follow_data) >= total:
                    break
                page += 1
                time.sleep(0.5)  # 礼貌延时，避免请求过快

            self._log(f"\n🎉 全部获取完成！共 {len(self.follow_data)} 条关注")

            self.root.after(0, lambda: self._set_ui_busy(False))
            if self.follow_data:
                self.root.after(0, lambda: self.export_btn.configure(state=tk.NORMAL))

        except requests.RequestException as e:
            self._log(f"❌ 网络错误: {e}")
            self.root.after(0, lambda: messagebox.showerror("网络错误", str(e)))
            self.root.after(0, lambda: self._set_ui_busy(False))
        except Exception as e:
            self._log(f"❌ 未知错误: {e}")
            self.root.after(0, lambda: messagebox.showerror("错误", str(e)))
            self.root.after(0, lambda: self._set_ui_busy(False))

    # ---------- 导出 Excel ----------

    def _export_excel(self):
        if not self.follow_data:
            messagebox.showinfo("提示", "没有数据可以导出，请先获取关注列表")
            return

        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
            initialfile=f"{self.my_name}_{self.my_uid}的B站关注列表.xlsx"
            if self.my_name
            else f"B站关注列表_{len(self.follow_data)}人.xlsx",
        )
        if not file_path:
            return

        threading.Thread(target=self._write_excel, args=(file_path,), daemon=True).start()

    def _write_excel(self, file_path: str):
        try:
            self.root.after(0, lambda: self.status_var.set("正在导出 Excel…"))
            self._log(f"正在导出 Excel 到: {os.path.basename(file_path)}")

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "关注列表"

            # 表头样式
            header_font = Font(name="Microsoft YaHei", bold=True, size=11, color="FFFFFF")
            header_fill = openpyxl.styles.PatternFill(
                start_color="00A1D6", end_color="00A1D6", fill_type="solid"
            )
            header_align = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

            # 写表头
            headers = ["序号", "UID", "用户名", "个人空间链接"]
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                cell.border = thin_border

            # 写数据
            data_font = Font(name="Microsoft YaHei", size=10)
            data_align = Alignment(horizontal="center", vertical="center")

            for idx, (uid, name) in enumerate(self.follow_data, 1):
                row = idx + 1
                ws.cell(row=row, column=1, value=idx).font = data_font
                ws.cell(row=row, column=1).alignment = data_align
                ws.cell(row=row, column=1).border = thin_border

                ws.cell(row=row, column=2, value=uid).font = data_font
                ws.cell(row=row, column=2).alignment = data_align
                ws.cell(row=row, column=2).border = thin_border

                ws.cell(row=row, column=3, value=name).font = data_font
                ws.cell(row=row, column=3).alignment = data_align
                ws.cell(row=row, column=3).border = thin_border

                link = f"https://space.bilibili.com/{uid}"
                cell = ws.cell(row=row, column=4, value=link)
                cell.font = Font(
                    name="Microsoft YaHei", size=10, color="0563C1", underline="single"
                )
                cell.alignment = data_align
                cell.border = thin_border
                cell.hyperlink = link

            # 设置列宽
            ws.column_dimensions["A"].width = 8
            ws.column_dimensions["B"].width = 14
            ws.column_dimensions["C"].width = 22
            ws.column_dimensions["D"].width = 40

            # 冻结首行
            ws.freeze_panes = "A2"

            wb.save(file_path)
            self._log(f"✅ 导出成功: {file_path}")
            self.root.after(
                0,
                lambda: messagebox.showinfo(
                    "导出成功",
                    f"已导出 {len(self.follow_data)} 条关注到:\n{file_path}",
                ),
            )
        except Exception as e:
            self._log(f"❌ 导出失败: {e}")
            self.root.after(0, lambda: messagebox.showerror("导出失败", str(e)))
        finally:
            self.root.after(0, lambda: self.status_var.set("就绪"))


def main():
    root = tk.Tk()
    app = BilibiliFollowExport(root)
    root.mainloop()


if __name__ == "__main__":
    main()
